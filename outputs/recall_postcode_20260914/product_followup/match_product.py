import json
import re
from pathlib import Path
from collections import defaultdict, Counter
from openpyxl import load_workbook

OUT=Path(__file__).resolve().parent
FILE=Path('C:/Users/Leo.Li/Downloads/Product_Registrations_Till_17072026.xlsx')
def clean(v):
    if v is None:return ''
    if isinstance(v,float) and v.is_integer():v=int(v)
    s=str(v).strip()
    return '' if s.lower() in {'','none','null','#','n/a','not assigned','-'} else s
def key(v):return re.sub('[^A-Z0-9]','',clean(v).upper())
def name(v):return ' '.join(re.findall('[A-Z0-9]+',clean(v).upper().replace('&',' AND ')))
def email(v):return clean(v).lower() if '@' in clean(v) else ''
def pc(v):
    s=clean(v)
    return s.zfill(4) if re.fullmatch(r'\d{3,4}',s) and int(s)>0 else ''
def compatible(a,b):
    aa=set(a.split())-{'AND'};bb=set(b.split())-{'AND'}
    return len(aa)>=2 and len(bb)>=2 and (aa<=bb or bb<=aa)
w=load_workbook(FILE,read_only=True,data_only=True)
records=[]
for s in w:
    for i,r in enumerate(s.iter_rows(min_row=2,values_only=True),2):
        records.append({'row':i,'sheet':s.title,'serial':key(r[3]),'vin':key(r[17]),'name':name(clean(r[7])+' '+clean(r[9])),'email':email(r[6]),'postcode':pc(r[13]),'raw_postcode':clean(r[13]),'street':clean(r[15]),'suburb':clean(r[16]),'country':clean(r[4]),'state':clean(r[14])})
w.close()
vehicles=defaultdict(set);emails=defaultdict(set);names=defaultdict(set)
for i,r in enumerate(records):
    for field,idx in [('serial',vehicles),('vin',vehicles),('email',emails),('name',names)]:
        if r[field]:idx[r[field]].add(i)
data=json.loads((OUT/'firebase_before.json').read_text(encoding='utf-8'))['tickets']
results=[]
for tid,t in sorted(data.items()):
    if clean(t.get('postcode')):continue
    raw=t.get('ticket',{});p=t.get('product',{});c=t.get('customer',{})
    ids={key(p.get('serialId')),key(p.get('chassisNumber')),key(raw.get('SerialID')),key(raw.get('ChassisNumber'))}-{''}
    em={email(c.get('email')),email(raw.get('ServiceRequesterEmail'))}-{''}
    ns={name(v.get('InvolvedPartyName')) for role,v in t.get('roles',{}).items() if role=='1001'}-{''}
    title=name(raw.get('TicketName'))
    for v in ids:title=re.sub(r'\b'+re.escape(v)+r'\b','',title)
    title=' '.join(title.split())
    if title:ns.add(title)
    vh=set().union(*(vehicles.get(v,set()) for v in ids))
    eh=set().union(*(emails.get(v,set()) for v in em))
    nh=set().union(*(names.get(v,set()) for v in ns))
    candidates=vh or eh
    selected={i for i in candidates if records[i]['postcode']}
    method='Exact vehicle ID'
    if not vh:
        method='Exact email and matching customer name'
        selected={i for i in selected if any(compatible(n,records[i]['name']) for n in ns)}
        if any(re.match(r'^(service|warranty|sales|customercare|admin|info)@',e) for e in em):selected=set()
    codes=sorted({records[i]['postcode'] for i in selected})
    status='matched' if len(codes)==1 else ('conflicting_postcodes' if codes else ('source_postcode_missing' if vh else ('unconfirmed_contact' if eh else ('name_only' if nh else 'no_match'))))
    # A one-character vehicle-ID difference is a review candidate, never a write.
    near=[]
    if status=='no_match':
        for vid in ids:
            if len(vid)<8:continue
            for candidate,indices in vehicles.items():
                if len(candidate)==len(vid) and sum(a!=b for a,b in zip(candidate,vid))==1:
                    for i in indices:
                        if records[i]['postcode'] and any(compatible(n,records[i]['name']) for n in ns):near.append(i)
    result={'ticket_id':tid,'status':status,'postcode':codes[0] if status=='matched' else '', 'method':method,'ticket_name':raw.get('TicketName',''),'ticket_ids':sorted(ids),'customer_names':sorted(ns),'customer_emails':sorted(em),'evidence':[dict(records[i],selected=i in selected,vehicle_match=i in vh,email_match=i in eh) for i in sorted(candidates|nh)],'near_vehicle_candidates':[records[i] for i in sorted(set(near))]}
    results.append(result)
(OUT/'match_results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'missing':len(results),'statuses':dict(Counter(r['status'] for r in results)),'matched':[{'ticket_id':r['ticket_id'],'postcode':r['postcode']} for r in results if r['status']=='matched']},indent=2))
for r in results:
    if r['status']!='no_match' or r['near_vehicle_candidates']:
        print(json.dumps(r,ensure_ascii=False))
