"""Audited recall postcode refresh: latest SAP first, then supplied fallback files."""
import copy
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

BASE = Path(__file__).resolve().parent / 'latest_20260914'
BASE.mkdir(exist_ok=True)
ROOT = Path(__file__).resolve().parents[2]
DOWNLOADS = Path('C:/Users/Leo.Li/Downloads')
FILES = {'SAP': 'SAPAnalyticsReport(Z578DB6F34E1CDB92C9D1CE) (1).xlsx',
         'Registration': 'Product_Registrations_14092026.xlsx', 'VCUST': 'VCUST.XLSX'}

def save(file, data):
    (BASE / file).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')

def read(file):
    return json.loads((BASE / file).read_text(encoding='utf-8'))

def clean(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): v = int(v)
    s = str(v).strip()
    return '' if s.lower() in {'', '#', 'none', 'null', 'n/a', 'not assigned', '-'} else s

def ident(v): return re.sub('[^A-Z0-9]', '', clean(v).upper())
def customer_id(v): return ident(v).lstrip('0')
def name(v): return ' '.join(re.findall('[A-Z0-9]+', clean(v).upper().replace('&', ' AND ')))
def email(v): return clean(v).lower() if '@' in clean(v) else ''
def phone(v):
    p = re.sub(r'\D', '', clean(v))
    if p.startswith('61') and len(p) == 11: p = '0' + p[2:]
    return p if len(p) >= 8 else ''
def postcode(v):
    s = clean(v)
    return s.zfill(4) if re.fullmatch(r'\d{3,4}', s) and int(s) else ''
def compatible(a, b):
    a, b = set(a.split()) - {'AND'}, set(b.split()) - {'AND'}
    return len(a) >= 2 and len(b) >= 2 and (a <= b or b <= a)

def ticket_keys(t):
    p, raw, c = t.get('product', {}), t.get('ticket', {}), t.get('customer', {})
    roles = [v for k, v in t.get('roles', {}).items() if k == '1001']
    ids = {ident(p.get('serialId')), ident(p.get('chassisNumber')), ident(raw.get('SerialID')), ident(raw.get('ChassisNumber'))} - {''}
    title = name(raw.get('TicketName'))
    for value in ids: title = re.sub(r'\b' + re.escape(value) + r'\b', '', title)
    values = {
        'vehicle': ids,
        'customer_id': {customer_id(r.get(k)) for r in roles for k in ['InvolvedPartyBusinessPartnerID', 'InvolvedPartyID']},
        'name': {name(r.get('InvolvedPartyName')) for r in roles} | {' '.join(title.split())},
        'email': {email(c.get('email')), email(raw.get('ServiceRequesterEmail'))},
        'phone': {phone(c.get(k)) for k in ['phone', 'mobile', 'serviceRequesterPhone']} | {phone(raw.get(k)) for k in ['ServiceRequesterPhone', 'ServiceRequesterMobile']},
    }
    return {k: sorted(v - {''}) for k, v in values.items()}

def firebase_ref():
    import firebase_admin
    from firebase_admin import credentials, db
    firebase_admin.initialize_app(credentials.Certificate(str(ROOT / 'firebase-service-account.json')),
        {'databaseURL': 'https://snowy-hr-report-default-rtdb.asia-southeast1.firebasedatabase.app', 'httpTimeout': 30})
    return db.reference('recallClaim')

def plan():
    from openpyxl import load_workbook
    records = []
    for source, filename in FILES.items():
        w = load_workbook(DOWNLOADS / filename, read_only=True, data_only=True)
        for s in w:
            header_row = 7 if source == 'SAP' else 1
            header = next(s.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
            assert header[19 if source == 'SAP' else 13 if source == 'Registration' else 12] == ('Postcode__c' if source == 'Registration' else 'Postal Code')
            for row, r in enumerate(s.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
                if source == 'SAP':
                    pc, vehicles, cust, ns, es, ps = r[19], [r[1], r[2]], [r[8]], [r[7]], [r[11]], [r[9], r[10]]
                elif source == 'Registration':
                    # SAP_ID__c is the registered-product ID, not the customer ID.
                    pc, vehicles, cust, ns, es, ps = r[13], [r[4], r[19]], [], [clean(r[8]) + ' ' + clean(r[10])], [r[7]], [r[11]]
                else:
                    pc, vehicles, cust, ns, es, ps = r[12], [], [r[0]], [r[6], r[9]], [r[35]], [r[16]]
                records.append({'source': source, 'file': filename, 'sheet': s.title, 'row': row, 'postcode': postcode(pc), 'raw_postcode': clean(pc),
                    'vehicle': sorted({ident(v) for v in vehicles} - {''}), 'customer_id': sorted({customer_id(v) for v in cust} - {''}),
                    'name': sorted({name(v) for v in ns} - {''}), 'email': sorted({email(v) for v in es} - {''}), 'phone': sorted({phone(v) for v in ps} - {''})})
        w.close()
    indexes = {kind: defaultdict(set) for kind in ['vehicle', 'customer_id', 'name', 'email', 'phone']}
    for i, r in enumerate(records):
        for kind, idx in indexes.items():
            for value in r[kind]: idx[value].add(i)

    def choose(keys, allowed):
        hits = {k: set().union(*(indexes[k].get(v, set()) for v in values)) & allowed for k, values in keys.items()}
        if hits['vehicle']:
            candidates, method = hits['vehicle'], 'Exact vehicle ID'
        elif hits['customer_id']:
            candidates, method = hits['customer_id'], 'Exact customer ID'
        else:
            candidates, method = hits['email'] | hits['phone'], 'Contact and customer name'
            candidates = {i for i in candidates if any(compatible(a, b) for a in keys['name'] for b in records[i]['name'])}
            if any(re.match(r'^(service|warranty|sales|customercare|admin|info)@', e) for e in keys['email']): candidates = set()
        valid = {i for i in candidates if records[i]['postcode']}
        codes = {records[i]['postcode'] for i in valid}
        if len(codes) > 1:
            # Disambiguate reused vehicles only with both exact contact and customer-name evidence.
            corroborated = {i for i in valid if i in hits['email'] | hits['phone'] and any(compatible(a,b) for a in keys['name'] for b in records[i]['name'])}
            if len({records[i]['postcode'] for i in corroborated}) == 1:
                valid = corroborated
                codes = {records[i]['postcode'] for i in valid}
                method += ' with matching contact and name'
        status = 'matched' if len(codes) == 1 else 'conflict' if len(codes) > 1 else 'missing_postcode' if candidates else 'no_match'
        return {'status': status, 'postcode': next(iter(codes)) if len(codes) == 1 else '', 'method': method,
                'evidence': [records[i] for i in sorted(candidates)], 'selected_rows': [(records[i]['source'],records[i]['row']) for i in sorted(valid)]}

    sap = {i for i,r in enumerate(records) if r['source'] == 'SAP'}
    fallback = set(range(len(records))) - sap
    results = []
    for tid, t in sorted(read('firebase_before.json')['tickets'].items()):
        assert t.get('ticketType') == 'Z011'
        keys = ticket_keys(t)
        primary = choose(keys, sap)
        secondary = None
        if primary['status'] in ['no_match', 'missing_postcode']:
            secondary = choose(keys, fallback)
        selected = secondary if secondary else primary
        results.append({'ticket_id': tid, 'keys': keys, 'existing_postcode': t.get('postcode'), 'source_group': 'fallback' if secondary else 'SAP', **selected, 'sap_status': primary['status']})
    save('match_results.json', results)
    summary = {'source_rows': dict(Counter(r['source'] for r in records)), 'total_tickets': len(results),
               'statuses': dict(Counter(r['status'] for r in results)),
               'new_to_fill': sum(r['status']=='matched' and not clean(r['existing_postcode']) for r in results),
               'changed_existing': sum(r['status']=='matched' and bool(clean(r['existing_postcode'])) and r['existing_postcode']!=r['postcode'] for r in results),
               'remaining_missing': [r['ticket_id'] for r in results if r['status']!='matched' and not clean(r['existing_postcode'])]}
    save('plan_summary.json',summary)
    print(json.dumps(summary,indent=2))
    print('conflicts', [(r['ticket_id'],r['source_group'],r['existing_postcode'],sorted({e['postcode'] for e in r['evidence']})) for r in results if r['status']=='conflict'])

def apply():
    ref = firebase_ref()
    before = ref.get()
    save('firebase_prewrite.json', before)
    results = read('match_results.json')
    edits = {r['ticket_id']: r for r in results if r['status'] == 'matched'}
    def update(current):
        result = copy.deepcopy(current)
        for tid, r in edits.items():
            t = result['tickets'].get(tid)
            assert t and t.get('ticketType')=='Z011' and ticket_keys(t)==r['keys'], 'Ticket changed during matching: '+tid
            assert t.get('postcode') in [r['existing_postcode'], r['postcode']], 'Postcode changed during matching: '+tid
            t['postcode'] = r['postcode']
        return result
    committed = ref.transaction(update)
    save('firebase_committed.json', committed)
    after = ref.get()
    save('firebase_after.json', after)
    failures = [tid for tid,r in edits.items() if after['tickets'].get(tid,{}).get('postcode')!=r['postcode']]
    missing = sorted(tid for tid,t in after['tickets'].items() if not clean(t.get('postcode')))
    summary = {'total_tickets':len(after['tickets']), 'missing_postcode_count':len(missing), 'missing_ticket_ids':missing,
               'verified_matched_tickets':len(edits)-len(failures),'failed_verification':failures,
               'other_fields_unchanged':update(before)==committed,'readback_equals_committed':after==committed,
               'completed_at_utc':datetime.now(timezone.utc).isoformat()}
    save('result_summary.json', summary)
    print(json.dumps(summary,indent=2))
    assert not failures and summary['other_fields_unchanged'] and summary['readback_equals_committed']

if __name__ == '__main__':
    if sys.argv[1] == 'snapshot':
        snap = firebase_ref().get()
        save('firebase_before.json', snap)
        print(json.dumps({'total':len(snap['tickets']),'missing':sum(not clean(t.get('postcode')) for t in snap['tickets'].values())}))
    elif sys.argv[1] == 'plan': plan()
    elif sys.argv[1] == 'apply': apply()
