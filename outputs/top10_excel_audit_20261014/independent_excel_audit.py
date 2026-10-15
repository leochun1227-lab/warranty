import json, collections, datetime
from decimal import Decimal
from pathlib import Path
import openpyxl

source = Path(r'C:/Users/Leo.Li/AppData/Local/Temp/ZF8C06456D7698BCB54F44D.xlsx')
w = openpyxl.load_workbook(source, read_only=True, data_only=True)
s = w.worksheets[0]
rows = list(s.iter_rows(min_row=11, values_only=True))
headers = list(next(s.iter_rows(min_row=10, max_row=10, values_only=True)))
approved_statuses = {'Sales Order Approved','Partially Picked','Dispatch Parts','Repair in Progress','Repairer Invoiced Received','Repairer Invoiced Processed','Approved Claims Closed (Closed)'}
def date(v):
    if isinstance(v, (datetime.datetime, datetime.date)): return v
    try: return datetime.datetime.fromisoformat(str(v))
    except ValueError: return None
def num(v):
    try: return Decimal(str(v))
    except Exception: return Decimal(0)
ids=collections.Counter(r[3] for r in rows)
eligible=[r for r in rows if r[25] in approved_statuses and date(r[19])]
missing=[r for r in rows if r[25] in approved_statuses and not date(r[19])]
def grouped(rs, period):
    c=collections.Counter((date(r[19]).strftime(period),r[4]) for r in rs)
    return [{'period':p,'type':t,'count':n} for (p,t),n in sorted(c.items())]
def amounts(rs):
    total=sum((num(r[29]) for r in rs),Decimal(0)); factory=sum((num(r[30]) for r in rs),Decimal(0)); labour=sum((num(r[31]) for r in rs),Decimal(0)); repairer=sum((num(r[32]) for r in rs),Decimal(0))
    diffs=[(r,num(r[29])-num(r[31])-num(r[30])-num(r[32])) for r in rs]
    nonzero=[(r,d) for r,d in diffs if abs(d)>Decimal('0.01')]
    return {'claim_total':float(total),'factory_parts':float(factory),'labour':float(labour),'repairer_parts_misc':float(repairer),'factory_plus_repairer':float(factory+repairer),'claim_minus_labour':float(total-labour),'difference':float(total-labour-factory-repairer),'rows_difference_gt_0_01':len(nonzero),'difference_samples':[{'excel_row':rows.index(r)+11,'ticket_id':r[3],'difference':float(d)} for r,d in nonzero[:8]]}
result = {
 'source':str(source),'sheet':s.title,'header_row':10,'last_row':s.max_row,'data_rows':len(rows),'headers':headers,
 'ticket_id_column':'D (B duplicates D; C is descriptive title despite Ticket ID header)',
 'unique_ticket_ids':len(ids),'duplicate_id_count':sum(n>1 for n in ids.values()),'extra_duplicate_id_rows':sum(n-1 for n in ids.values()),'exact_duplicate_rows':len(rows)-len(set(rows)), 'B_D_mismatches':sum(r[1]!=r[3] for r in rows),
 'statuses':dict(collections.Counter(r[25] for r in rows)), 'types':dict(collections.Counter(r[4] for r in rows)),
 'approved_statuses':sorted(approved_statuses),'eligible_count':len(eligible),'eligible_by_year_type':grouped(eligible,'%Y'),
 'eligible_2026_by_month_type':grouped([r for r in eligible if date(r[19]).year==2026],'%Y-%m'),
 'approved_status_missing_date_count':len(missing),'approved_status_missing_date_by_status':dict(collections.Counter(r[25] for r in missing)),
 'approved_status_missing_date_created_2026':sum(bool(date(r[14]) and date(r[14]).year==2026) for r in missing),
 'valid_approval_date_noneligible_statuses':dict(collections.Counter(r[25] for r in rows if date(r[19]) and r[25] not in approved_statuses)),
 'amounts_all':amounts(rows),'amounts_eligible':amounts(eligible),'amounts_eligible_2026':amounts([r for r in eligible if date(r[19]).year==2026]),
 'issue_position_columns':[h for h in headers if 'issue' in str(h).lower() or 'position' in str(h).lower()],
 'approval_date_range':[str(min(date(r[19]) for r in eligible)),str(max(date(r[19]) for r in eligible))]
}
dest=Path('outputs/top10_excel_audit_20261014/independent_excel.json')
dest.write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result,indent=2))
